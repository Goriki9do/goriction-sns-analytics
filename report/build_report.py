"""
YouTube/TikTok/Instagram/Xの各exports CSVを1つのExcelにまとめる。

各SNSのfetchスクリプトを実行した後にこれを実行すると、
report/goriction_sns_report.xlsx が(既存があれば上書きで)作り直される。
VBAマクロは使わず、毎回Pythonでゼロから作り直す方式にしているので、
Excel起動時のマクロ有効化の警告も出ない。

使い方:
    python build_report.py
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
OUTPUT_PATH = Path(__file__).parent / "goriction_sns_report.xlsx"

def to_naive_jst(series: pd.Series) -> pd.Series:
    """Excelはtz付きdatetimeを書けないので、JSTに変換してtz情報を外す。"""
    return series.dt.tz_convert("Asia/Tokyo").dt.tz_localize(None)


COLUMNS = [
    "platform",
    "posted_at",
    "caption",
    "likes",
    "views_or_impressions",
    "comments_raw",
    "unique_commenters",
    "permalink",
]


def load_youtube() -> pd.DataFrame:
    paths = list((ROOT / "youtube" / "exports_local").glob("*/videos.csv"))
    if not paths:
        print("youtube: exports_local/*/videos.csv が見つからんかった。スキップするで。")
        return pd.DataFrame(columns=COLUMNS)
    df = pd.concat([pd.read_csv(p) for p in paths], ignore_index=True)
    out = pd.DataFrame(
        {
            "platform": "YouTube",
            "posted_at": to_naive_jst(pd.to_datetime(df["published_at"], errors="coerce", utc=True)),
            "caption": df["title"],
            "likes": df["likes"],
            "views_or_impressions": df["views"],
            "comments_raw": df["comment_count_raw"],
            "unique_commenters": df["unique_commenters"],
            "permalink": "https://www.youtube.com/watch?v=" + df["video_id"].astype(str),
        }
    )
    return out


def load_tiktok() -> pd.DataFrame:
    path = ROOT / "tiktok" / "exports" / "videos.csv"
    if not path.exists():
        print("tiktok: exports/videos.csv が見つからんかった。スキップするで。")
        return pd.DataFrame(columns=COLUMNS)
    df = pd.read_csv(path)
    # 同じ動画が複数回fetchされていることがあるので、video_urlで最新のfetched_atだけ残す
    df = df.sort_values("fetched_at").drop_duplicates(subset="video_url", keep="last")
    out = pd.DataFrame(
        {
            "platform": "TikTok",
            # TikTok Studioの日付表示には年が含まれない(例:「9月19日」)。
            # 年を跨ぐ蓄積時に誤認しないよう、日付として解釈せず文字列のまま保持する。
            "posted_at": df["post_date"].astype(str) + "(年不明)",
            "caption": df["title"],
            "likes": df["likes"],
            "views_or_impressions": df["views"],
            "comments_raw": df["comment_count_raw"],
            "unique_commenters": df["unique_commenters"],
            "permalink": df["video_url"],
        }
    )
    return out


def load_instagram() -> pd.DataFrame:
    path = ROOT / "instagram" / "exports" / "media.csv"
    if not path.exists():
        print("instagram: exports/media.csv が見つからんかった。スキップするで。")
        return pd.DataFrame(columns=COLUMNS)
    df = pd.read_csv(path)
    out = pd.DataFrame(
        {
            "platform": "Instagram",
            "posted_at": to_naive_jst(pd.to_datetime(df["timestamp"], errors="coerce", utc=True)),
            "caption": df["caption"],
            "likes": df["like_count"],
            "views_or_impressions": pd.NA,  # インサイト未実装のため表示回数は取得できていない
            "comments_raw": df["comment_count_raw"],
            "unique_commenters": df["unique_commenters"],
            "permalink": df["permalink"],
        }
    )
    return out


def load_x() -> pd.DataFrame:
    path = ROOT / "x" / "exports" / "posts.csv"
    if not path.exists():
        print("x: exports/posts.csv が見つからんかった。スキップするで。")
        return pd.DataFrame(columns=COLUMNS)
    df = pd.read_csv(path)
    out = pd.DataFrame(
        {
            "platform": "X",
            "posted_at": to_naive_jst(pd.to_datetime(df["created_at"], errors="coerce", utc=True)),
            "caption": df["text"],
            "likes": df["like_count"],
            "views_or_impressions": df["impression_count"],
            "comments_raw": df["reply_count_raw"],
            "unique_commenters": df["unique_commenters"],
            "permalink": df["permalink"],
        }
    )
    return out


def build_summary(all_posts: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for platform in ["YouTube", "TikTok", "Instagram", "X"]:
        sub = all_posts[all_posts["platform"] == platform]
        if sub.empty:
            continue
        rows.append(
            {
                "platform": platform,
                "投稿数": len(sub),
                "合計いいね": int(sub["likes"].fillna(0).sum()),
                "合計再生/表示回数": (
                    int(sub["views_or_impressions"].fillna(0).sum())
                    if sub["views_or_impressions"].notna().any()
                    else None
                ),
                "合計コメント/リプライ(raw)": int(sub["comments_raw"].fillna(0).sum()),
            }
        )
    return pd.DataFrame(rows)


def autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 60)


def add_bar_chart(ws, summary_row_count: int, value_col: int, title: str, anchor: str) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    # カテゴリ(プラットフォーム)ごとに自動で色分けされるので、
    # どれがどの色か分かるように凡例は残す
    chart.legend.position = "b"

    data = Reference(ws, min_col=value_col, min_row=1, max_row=summary_row_count + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=summary_row_count + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    # 一番高いバーの数値ラベルがグラフタイトルと重ならないよう、上に余白を確保する
    values = [
        ws.cell(row=r, column=value_col).value
        for r in range(2, summary_row_count + 2)
    ]
    values = [v for v in values if isinstance(v, (int, float))]
    if values:
        chart.y_axis.scaling.max = max(values) * 1.2

    series = chart.series[0]
    series.varyColors = True
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    series.dLbls.showCatName = False
    series.dLbls.showSerName = False
    series.dLbls.showLegendKey = False
    series.dLbls.numFmt = "#,##0"

    chart.height = 8
    chart.width = 13
    ws.add_chart(chart, anchor)


def main() -> None:
    frames = [load_youtube(), load_tiktok(), load_instagram(), load_x()]
    all_posts = pd.concat(frames, ignore_index=True)[COLUMNS]

    if all_posts.empty:
        print("どのSNSのデータも見つからんかった。各SNSのfetchスクリプトを先に実行してな。")
        sys.exit(1)

    summary = build_summary(all_posts)

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="サマリー", index=False)

        for platform in ["YouTube", "TikTok", "Instagram", "X"]:
            sub = all_posts[all_posts["platform"] == platform].copy()
            if sub.empty:
                continue
            sub = sub.sort_values("likes", ascending=False)
            sub.drop(columns=["platform"]).to_excel(writer, sheet_name=platform, index=False)

        notes = pd.DataFrame(
            {
                "注意点": [
                    "Instagramは現状いいね・コメント数のみ取得。表示回数/リーチ(インサイト)は未実装。",
                    "Xのリプライ取得は直近7日以内が対象。7日より前の投稿はunique_commentersが過小(0)になりうる。",
                    "TikTokの投稿日は年情報が無いため文字列のまま保持している(年跨ぎで誤認しないよう注意)。",
                    "YouTube/TikTokの再生回数とXのインプレッション(表示回数)は測定定義が異なる目安値。",
                    "このExcelはfetch時点のスナップショット。日々の推移を追うには過去分を別途残す必要がある。",
                ]
            }
        )
        notes.to_excel(writer, sheet_name="サマリー", index=False, startrow=len(summary) + 3)

    from openpyxl import load_workbook

    wb = load_workbook(OUTPUT_PATH)
    ws = wb["サマリー"]
    ws["A1"].font = Font(bold=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    autosize_columns(ws)

    chart_start_row = len(summary) + len(notes) + 6
    add_bar_chart(ws, len(summary), value_col=2, title="投稿数", anchor=f"A{chart_start_row}")
    add_bar_chart(ws, len(summary), value_col=3, title="合計いいね", anchor=f"A{chart_start_row + 18}")
    add_bar_chart(ws, len(summary), value_col=4, title="合計再生/表示回数", anchor=f"A{chart_start_row + 36}")

    for platform in ["YouTube", "TikTok", "Instagram", "X"]:
        if platform in wb.sheetnames:
            autosize_columns(wb[platform])

    wb.save(OUTPUT_PATH)
    print(f"完了。{OUTPUT_PATH} に保存したで。")


if __name__ == "__main__":
    main()
